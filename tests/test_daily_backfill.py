from datetime import date
from pathlib import Path

import openpyxl
import pytest

from src.core.file_naming import resolve_file_date
from src.services.daily_file_service import backfill_dailies, find_missing_dates

_NAME = "리와인드 중산점 데일리 {m}.{d}.xlsx"


@pytest.fixture
def daily_dir(tmp_path, monkeypatch):
    """'데일리매출'·'데일리지출' 시트를 가진 최소한의 데일리 파일 1개를 만든다."""
    monkeypatch.setattr(
        "src.services.daily_file_service.get_expense_daily_sheet",
        lambda: "데일리지출",
    )

    def make(month: int, day: int) -> Path:
        wb = openpyxl.Workbook()
        sales = wb.active
        sales.title = "데일리매출"
        sales.cell(12, 4).value = "홍길동"   # D: 회원명 (센터)
        sales.cell(12, 6).value = 100000     # F: 금액
        sales.cell(5, 13).value = "=SUM(1)"  # M5: 수식은 보존되어야 함
        expense = wb.create_sheet("데일리지출")
        expense.cell(6, 3).value = "10/10"   # C: 일자
        expense.cell(6, 5).value = 5000
        path = tmp_path / _NAME.format(m=month, d=day)
        wb.save(str(path))
        return path

    return make


class TestResolveFileDate:
    def test_same_year(self):
        assert resolve_file_date("데일리 10.11.xlsx", date(2026, 10, 13)) == date(2026, 10, 11)

    def test_new_year_boundary(self):
        # 1월 2일에 본 '12.31'은 올해가 아니라 작년
        assert resolve_file_date("데일리 12.31.xlsx", date(2026, 1, 2)) == date(2025, 12, 31)


class TestFindMissingDates:
    def test_lists_only_absent_days(self, daily_dir):
        source = daily_dir(10, 11)
        # 10.12는 이미 있다고 가정 (전날 '내일 파일 미리 생성'이 만든 파일)
        (source.parent / _NAME.format(m=10, d=12)).write_bytes(b"")

        missing = find_missing_dates(source, date(2026, 10, 14))
        assert missing == [date(2026, 10, 13), date(2026, 10, 14)]


class TestBackfillDailies:
    def test_fills_whole_gap(self, daily_dir):
        source = daily_dir(10, 11)

        created = backfill_dailies(source, date(2026, 10, 14))

        assert [p.name for p in created] == [
            _NAME.format(m=10, d=12),
            _NAME.format(m=10, d=13),
            _NAME.format(m=10, d=14),
        ]
        assert all(p.exists() for p in created)

    def test_created_files_keep_form_but_no_sales(self, daily_dir):
        source = daily_dir(10, 11)

        created = backfill_dailies(source, date(2026, 10, 13))

        wb = openpyxl.load_workbook(str(created[0]))
        assert wb.sheetnames == ["데일리매출", "데일리지출"]
        assert wb["데일리매출"].cell(12, 4).value is None   # 매출 입력은 비움
        assert wb["데일리매출"].cell(12, 6).value is None
        assert wb["데일리매출"].cell(5, 13).value == "=SUM(1)"  # 수식은 보존
        assert wb["데일리지출"].cell(6, 3).value is None

    def test_skips_existing_files(self, daily_dir):
        source = daily_dir(10, 11)
        existing = source.parent / _NAME.format(m=10, d=12)
        existing.write_bytes(b"sentinel")

        created = backfill_dailies(source, date(2026, 10, 13))

        assert [p.name for p in created] == [_NAME.format(m=10, d=13)]
        assert existing.read_bytes() == b"sentinel"  # 기존 파일은 덮어쓰지 않음

    def test_no_gap_returns_empty(self, daily_dir):
        source = daily_dir(10, 13)
        assert backfill_dailies(source, date(2026, 10, 13)) == []

    def test_rejects_gap_over_limit(self, daily_dir):
        source = daily_dir(10, 11)
        with pytest.raises(ValueError, match="직접 선택"):
            backfill_dailies(source, date(2026, 10, 13), max_days=1)

    def test_crosses_month_boundary(self, daily_dir):
        source = daily_dir(9, 30)

        created = backfill_dailies(source, date(2026, 10, 2))

        assert [p.name for p in created] == [
            _NAME.format(m=10, d=1),
            _NAME.format(m=10, d=2),
        ]
