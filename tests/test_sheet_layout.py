"""매출 시트 열 배치 — 틀리면 매출이 엉뚱한 칸에 기록되므로 열 단위로 고정해둔다."""
from datetime import date

import pytest

from src.services.entry_service import (
    CURRENT_LAYOUT,
    LEGACY_LAYOUT,
    PaymentEntry,
    _col_letter,
    _offsets,
    _vat_formula,
)
from src.services.total_sales_service import find_monthly_file_sheet_name


def _columns(layout, section):
    start = layout.section_start[section]
    return {
        key: (_col_letter(start + off) if off >= 0 else None)
        for key, off in _offsets(layout, section).items()
    }


class TestCurrentLayout:
    """데일리 · 월별 총매 — 센터 종목과 회원권 사이에 신/재 칸이 있다."""

    def test_center_columns(self):
        assert _columns(CURRENT_LAYOUT, "센터") == {
            "date": "B", "day": "C", "name": "D", "category": "E",
            "membership_type": "F", "membership": "G", "amount": "H",
            "vat": "I", "payment": "J", "approval": "K", "fc": "L",
            "manager": "M", "note": "N",
        }

    def test_lesson_shifts_right_but_keeps_12_columns(self):
        cols = _columns(CURRENT_LAYOUT, "레슨")
        assert cols["date"] == "Q"          # 센터가 한 칸 늘어 P → Q
        assert cols["membership_type"] is None
        assert cols["membership"] == "U"
        assert cols["note"] == "AB"


class TestLegacyLayout:
    """전체 총매 — 신/재 칸이 없는 기존 양식."""

    def test_center_columns(self):
        cols = _columns(LEGACY_LAYOUT, "센터")
        assert cols["membership_type"] is None
        assert cols["membership"] == "F"
        assert cols["amount"] == "G"
        assert cols["note"] == "M"

    def test_lesson_starts_at_p(self):
        assert _columns(LEGACY_LAYOUT, "레슨")["date"] == "P"


class TestVatFormula:
    @pytest.mark.parametrize(
        ("layout", "section", "amount_cell", "payment_cell"),
        [
            (CURRENT_LAYOUT, "센터", "H12", "J12"),
            (CURRENT_LAYOUT, "레슨", "V12", "X12"),
            (LEGACY_LAYOUT, "센터", "G12", "I12"),
            (LEGACY_LAYOUT, "레슨", "U12", "W12"),
        ],
    )
    def test_references_shift_with_layout(self, layout, section, amount_cell, payment_cell):
        start = layout.section_start[section]
        formula = _vat_formula(section, start, 12, _offsets(layout, section))
        assert amount_cell in formula
        assert payment_cell in formula


class TestPaymentEntry:
    def test_membership_type_defaults_empty(self):
        entry = PaymentEntry(
            entry_date=date(2026, 10, 23), name="홍길동", category="헬스",
            membership="3개월", amount=300000, payment_method="카드",
        )
        assert entry.membership_type == ""


class TestMonthlyFileSheetName:
    @pytest.mark.parametrize(
        ("year", "month", "kind", "expected"),
        [
            (2026, 10, "매출", "26.10월매출"),
            (2026, 10, "지출", "26.10월지출"),
            (2026, 9, "매출", "26.9월매출"),
        ],
    )
    def test_finds_sheet(self, year, month, kind, expected):
        names = ["표지", "26.9월매출", "26.10월매출", "26.10월지출"]
        assert find_monthly_file_sheet_name(names, year, month, kind) == expected

    def test_sales_lookup_does_not_match_expense_sheet(self):
        with pytest.raises(ValueError):
            find_monthly_file_sheet_name(["26.10월지출"], 2026, 10, "매출")

    def test_missing_month_raises_with_sheet_list(self):
        with pytest.raises(ValueError, match="26.11월매출"):
            find_monthly_file_sheet_name(["26.10월매출"], 2026, 11, "매출")


class TestYearlyTotalFile:
    """연말 — 전체 총매 파일을 새 연도로 넘긴다."""

    def test_builds_next_year_name(self):
        from src.services.yearly_total_service import build_year_path

        assert build_year_path("/x/리와인드 총매출 2026.xlsx", 2027).name == \
            "리와인드 총매출 2027.xlsx"

    def test_reads_year_from_name(self):
        from src.services.yearly_total_service import file_year

        assert file_year("/x/리와인드 총매출 2026.xlsx") == 2026
        assert file_year("/x/총매출.xlsx") is None

    def test_name_without_year_explains_what_to_do(self):
        from src.services.yearly_total_service import build_year_path

        with pytest.raises(ValueError, match="직접 만들어"):
            build_year_path("/x/총매출.xlsx", 2027)

    def test_only_old_month_sheets_are_removed(self):
        from src.services.yearly_total_service import _old_month_sheet_names

        names = [
            "표지", "연간요약",
            "총 매출26년 11월", "총 지출26년 11월",
            "총 매출26년 12월", "총 지출26년 12월",
            "총 매출27년 1월", "총 지출27년 1월",
        ]
        removed = _old_month_sheet_names(names, 2027)

        assert removed == [
            "총 매출26년 11월", "총 지출26년 11월",
            "총 매출26년 12월", "총 지출26년 12월",
        ]
        assert "표지" not in removed        # 월 시트가 아닌 것은 건드리지 않는다
        assert "연간요약" not in removed


class TestReaderMatchesWriterLayout:
    """입력·조회가 서로 다른 열 배치를 들고 있어 내역 조회가 비던 적이 있다."""

    def _daily(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "데일리매출"
        # 센터: 신/재(F) 포함 13칸
        for cell, value in [
            ("C12", 24), ("D12", "홍길동"), ("E12", "헬스"), ("F12", "신규"),
            ("G12", "헬스 3개월"), ("H12", 300000), ("J12", "카드(VAN)"), ("L12", "정준"),
        ]:
            ws[cell] = value
        # 레슨: 신/재 없이 Q열부터
        for cell, value in [
            ("R13", 24), ("S13", "김철수"), ("T13", "PT"),
            ("U13", "PT 10회"), ("V13", 900000), ("X13", "현금"),
        ]:
            ws[cell] = value
        path = tmp_path / "리와인드 데일리 10.24.xlsx"
        wb.save(str(path))
        return path

    def test_reads_both_sections_of_new_layout(self, tmp_path):
        from src.services.entry_reader_service import read_sales_entries

        rows = read_sales_entries(self._daily(tmp_path))

        assert len(rows) == 2
        center = next(r for r in rows if r.section == "센터")
        assert (center.name, center.membership_type, center.membership) == \
            ("홍길동", "신규", "헬스 3개월")
        assert center.amount == 300000      # 신/재 칸만큼 밀린 금액 열을 읽는다
        assert center.fc == "정준"

        lesson = next(r for r in rows if r.section == "레슨")
        assert (lesson.name, lesson.membership, lesson.amount) == ("김철수", "PT 10회", 900000)
        assert lesson.membership_type == ""  # 레슨에는 신/재 칸이 없다

    def test_text_in_amount_column_does_not_kill_the_whole_list(self, tmp_path):
        """양식이 또 어긋나도 조회가 통째로 비지 않고 읽히는 만큼은 보여준다."""
        import openpyxl

        from src.services.entry_reader_service import read_sales_entries

        path = self._daily(tmp_path)
        wb = openpyxl.load_workbook(str(path))
        wb["데일리매출"]["H12"] = "삼십만원"   # 금액 칸에 글자
        wb.save(str(path))

        rows = read_sales_entries(path)

        assert len(rows) == 2
        assert next(r for r in rows if r.section == "센터").amount == 0
